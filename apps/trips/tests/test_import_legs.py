"""Phase D of docs/plans/trip-items-and-confirmations.md: an imported itinerary's travel
tables start as legs (import_itinerary runs travel_tables.convert_automatically), and the
section editor shows a travel section's legs."""
import io
import tempfile
from contextlib import redirect_stderr, redirect_stdout
from datetime import date
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse

from apps.trips.models import Day, Section, TransportLeg, Trip, TripGrant, TripRole
from apps.trips.tests.test_docx_europe import _table, build_europe_document

User = get_user_model()


def run_import(path, **kwargs):
    out, err = io.StringIO(), io.StringIO()
    with redirect_stdout(out), redirect_stderr(err):
        call_command("import_itinerary", str(path), **kwargs)
    return out.getvalue()


def build_taos_document(tmpdir):
    """A Taos-shaped day: a flights table and a schedule table naming the same hops."""
    import docx

    path = Path(tmpdir) / "taos.docx"
    document = docx.Document()
    document.add_paragraph("Mini Taos — Testville, New Mexico")
    document.add_paragraph("September 12–13, 2026")
    document.add_paragraph("Day 1 — Saturday, September 12 — Fly home")
    _table(document, ("Flight", "Route", "Times", "Seats"),
           [("DL 3726", "ABQ → SLC", "Depart 6:10 PM / Arrive 8:02 PM", "12A")])
    _table(document, ("Time", "Stage", "Details"), [
        ("~3:00 PM", "Drive to ABQ Airport", "~15–20 min"),  # "possible": the word "to" only
        ("6:10 PM", "Depart ABQ → SLC", "DL 3726"),
        ("8:02 PM", "Arrive SLC", "Layover")])
    document.add_paragraph("Day 2 — Sunday, September 13 — Home")
    document.save(path)
    return path


class ImportStartsAsLegsTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def test_europe_hops_become_planned_legs_in_a_travel_section(self):
        path = build_europe_document(self.tmp.name)
        report = run_import(path)  # dry run
        self.assertIn("Travel rows -> legs", report)
        self.assertIn("Butzbach -> Frankfurt Hbf", report)
        self.assertFalse(Trip.objects.exists())

        report = run_import(path, apply=True)
        self.assertIn("made 1 planned legs", report)
        leg = TransportLeg.objects.get()
        self.assertEqual((leg.origin, leg.destination, leg.is_planned, leg.est_duration_str, leg.operator),
                         ("Butzbach", "Frankfurt Hbf", True, "~45 min", "RE30"))
        self.assertEqual(leg.section.section_type, Section.Type.TRAVEL)
        self.assertEqual(leg.section.day.date, date(2027, 9, 17))
        self.assertEqual(leg.date, date(2027, 9, 17))

    def test_keep_tables(self):
        run_import(build_europe_document(self.tmp.name), apply=True, keep_tables=True)
        self.assertFalse(TransportLeg.objects.exists())
        self.assertTrue(Section.objects.filter(section_type="logistics_table").exists())
        self.assertFalse(Section.objects.filter(section_type="travel").exists())

    def test_replace_reimports_cleanly(self):
        path = build_europe_document(self.tmp.name)
        run_import(path, apply=True)
        run_import(path, apply=True, replace=True)
        self.assertEqual(Trip.objects.count(), 1)
        self.assertEqual(TransportLeg.objects.count(), 1)

    def test_taos_schedule_links_to_the_documents_flight(self):
        run_import(build_taos_document(self.tmp.name), apply=True)
        flight = TransportLeg.objects.get()  # linked, not doubled
        self.assertEqual((flight.origin, flight.destination, flight.is_planned), ("ABQ", "SLC", False))
        section = flight.section
        self.assertEqual(section.section_type, Section.Type.TRAVEL)
        # The "possible" drive row and the arrival row stay as reference rows.
        self.assertEqual([row[1] for row in section.content["rows"]], ["Drive to ABQ Airport", "Arrive SLC"])


class SectionEditorTests(TestCase):
    def setUp(self):
        self.editor = User.objects.create_user("ed", "ed@example.com", "pw-1234-abcd")
        self.trip = Trip.objects.create(name="Europe", start_date=date(2027, 9, 20), end_date=date(2027, 9, 21),
                                        created_by=self.editor)
        TripGrant.objects.create(trip=self.trip, user=self.editor, role=TripRole.EDITOR)
        self.day = Day.objects.create(trip=self.trip, day_number=1, date=date(2027, 9, 20))
        self.client.force_login(self.editor)

    def test_travel_section_lists_its_legs_with_edit_links(self):
        section = Section.objects.create(day=self.day, section_type="travel", title="Train south", content={})
        leg = TransportLeg.objects.create(trip=self.trip, date=self.day.date, origin="Butzbach",
                                          destination="Frankfurt Hbf", section=section)
        url = reverse("trips:section_edit", args=[self.trip.pk, self.day.pk, section.pk])
        page = self.client.get(url)
        self.assertContains(page, "Legs shown here")
        self.assertContains(page, reverse("trips:leg_edit", args=[self.trip.pk, leg.pk]) + "?next=" + url)
        self.assertContains(page, reverse("trips:leg_quick_add", args=[self.trip.pk, self.day.pk]))
        added = self.client.post(reverse("trips:leg_quick_add", args=[self.trip.pk, self.day.pk]), {
            "origin": "Frankfurt Hbf", "destination": "Augsburg Hbf", "mode": "train", "est_duration": "2.5 hrs",
            "section": section.pk, "next": url})
        self.assertRedirects(added, url, fetch_redirect_response=False)  # back to the section page
        self.assertEqual(section.legs.count(), 2)

    def test_logistics_table_with_arrow_rows_offers_turn_into_legs(self):
        section = Section.objects.create(day=self.day, section_type="logistics_table", title="Trains", content={
            "columns": ["Leg", "Time"], "rows": [["Butzbach → Frankfurt Hbf", "~45 min"]]})
        page = self.client.get(reverse("trips:section_edit", args=[self.trip.pk, self.day.pk, section.pk]))
        self.assertContains(page, reverse("trips:section_to_legs", args=[self.trip.pk, self.day.pk, section.pk]))
