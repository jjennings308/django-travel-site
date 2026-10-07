"""Tests for the DOCX importer.

The documents themselves are never committed, so the parser is exercised
against a small synthetic document built with python-docx. That keeps the
tests hermetic while still covering the parts that are easy to get wrong:
the euro/dollar price fallback, meals that exist only in the Food at a
Glance table, and confirmations written as prose.
"""

import io
import tempfile
from contextlib import redirect_stdout, redirect_stderr
from pathlib import Path

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase

from trips.docx_import import _compact_confirmation, parse_document
from trips.models import Day, Meal, Trip


def run_import(path, **kwargs):
    """Call the command with its report output swallowed.

    The report is the point of the command, but it is not what these tests are
    asserting on, and letting it print would bury the test results.
    """
    out, err = io.StringIO(), io.StringIO()
    with redirect_stdout(out), redirect_stderr(err):
        call_command("import_itinerary", str(path), **kwargs)


def build_document(tmpdir):
    """Write a miniature itinerary with the shapes that broke in real files."""
    import docx

    path = Path(tmpdir) / "mini.docx"
    document = docx.Document()

    document.add_paragraph("Mini Trip — Testville, New Mexico")
    document.add_paragraph("September 12–13, 2026")

    overview = document.add_table(rows=1, cols=2)
    overview.rows[0].cells[0].text = "Item"
    overview.rows[0].cells[1].text = "Details"
    for label, value in (("Travelers", "Ada & Bo Lovelace"), ("Home airport", "BWI")):
        row = overview.add_row()
        row.cells[0].text = label
        row.cells[1].text = value

    document.add_paragraph("Day 1 — Saturday, September 12 — Arrival")
    document.add_paragraph("Phase 1 — Getting There")

    logistics = document.add_table(rows=1, cols=3)
    for i, head in enumerate(("Time", "Stop", "Details")):
        logistics.rows[0].cells[i].text = head
    row = logistics.add_row()
    row.cells[0].text = "~12:00 PM"
    row.cells[1].text = "Lunch: Corner Cafe"
    row.cells[2].text = "12 Main St — good tacos"

    restaurants = document.add_table(rows=1, cols=3)
    for i, head in enumerate(("Restaurant", "Price", "Why")):
        restaurants.rows[0].cells[i].text = head
    row = restaurants.add_row()
    row.cells[0].text = "Fancy Grill"
    row.cells[1].text = "€€€"  # the stray euro sign seen in the real document
    row.cells[2].text = "Steakhouse worth the splurge"

    document.add_paragraph("Food at a Glance")
    glance = document.add_table(rows=1, cols=6)
    for i, head in enumerate(("Day", "Meal", "Restaurant", "Cuisine", "Price", "Atmosphere")):
        glance.rows[0].cells[i].text = head
    for cells in (
        ("9/12", "Lunch", "Corner Cafe", "Mexican", "$$", "Casual, quick"),
        ("9/12", "Dinner", "Fancy Grill", "Steakhouse", "$$$", "Upscale"),
    ):
        row = glance.add_row()
        for i, value in enumerate(cells):
            row.cells[i].text = value

    document.add_paragraph("Quick Reference")
    confirmations = document.add_table(rows=1, cols=2)
    confirmations.rows[0].cells[0].text = "Item"
    confirmations.rows[0].cells[1].text = "Confirmation #"
    row = confirmations.add_row()
    row.cells[0].text = "Hotel"
    row.cells[1].text = "Booking.com Confirmation 1234567 · PIN 9999 (confidential)"

    document.save(path)
    return path


class CompactConfirmationTests(TestCase):
    def test_extracts_code_and_drops_confidential(self):
        code, remainder = _compact_confirmation(
            "Booking.com Confirmation 5305434791 · PIN 7553 (confidential)"
        )
        self.assertEqual(code, "5305434791")
        self.assertNotIn("confidential", remainder)

    def test_keeps_hyphenated_code_whole(self):
        code, _ = _compact_confirmation("Reservation-1216125 · EVRN0002048844001")
        self.assertEqual(code, "Reservation-1216125")

    def test_short_value_is_untouched(self):
        code, remainder = _compact_confirmation("GCH7JS")
        self.assertEqual(code, "GCH7JS")
        self.assertEqual(remainder, "")


class ParseDocumentTests(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls._tmp = tempfile.TemporaryDirectory()
        cls.path = build_document(cls._tmp.name)
        cls.parsed = parse_document(cls.path)

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()
        super().tearDownClass()

    def test_trip_basics(self):
        self.assertEqual(self.parsed.trip["name"], "Mini Trip — Testville, New Mexico")
        # The parser hands back ISO strings; the command converts them on write.
        self.assertEqual(self.parsed.trip["start_date"], "2026-09-12")
        self.assertEqual(self.parsed.trip["end_date"], "2026-09-13")

    def test_travelers_get_shared_surname(self):
        self.assertEqual(
            [t["name"] for t in self.parsed.travelers],
            ["Ada Lovelace", "Bo Lovelace"],
        )

    def test_euro_price_falls_back_to_summary(self):
        meal = next(m for m in self.parsed.days[0]["meals"] if m["name"] == "Fancy Grill")
        self.assertEqual(meal["price_range"], "$$$")
        self.assertTrue(any("euro" in w or "not one of" in w for w in self.parsed.warnings))

    def test_meal_only_in_summary_is_still_created(self):
        # "Corner Cafe" appears in a logistics row, not a Restaurant|Price|Why
        # table, so it would be lost without the Food at a Glance back-fill.
        meal = next(m for m in self.parsed.days[0]["meals"] if m["name"] == "Corner Cafe")
        self.assertEqual(meal["meal_type"], "lunch")
        self.assertEqual(meal["cuisine"], "Mexican")
        self.assertEqual(meal["price_range"], "$$")
        self.assertEqual(meal["why_recommended"], "Casual, quick")
        self.assertIn("12 Main St", meal["reservation_notes"])

    def test_all_meals_land_on_the_right_day(self):
        self.assertEqual(len(self.parsed.days[0]["meals"]), 2)

    def test_confirmation_is_compacted(self):
        self.assertEqual(
            self.parsed.confirmations[0]["confirmation_number"], "1234567"
        )


class StayDateTests(TestCase):
    def test_checkout_row_without_a_year_uses_the_checkin_year(self):
        from trips.docx_import import _parse_stay_dates

        fields = {
            "check-in": "Sept 13, 4:00 PM",
            "check-out": "Sept 20, 11:00 AM",
        }
        check_in, check_out = _parse_stay_dates(fields, "2026-09-13", None)
        self.assertEqual(check_in, "2026-09-13")
        self.assertEqual(check_out, "2026-09-20")

    def test_stay_row_shape(self):
        from trips.docx_import import _parse_stay_dates

        fields = {"stay": "1 night, Sept 12 — check-in 4:00 PM, check-out Sept 13, 11:00 AM"}
        check_in, check_out = _parse_stay_dates(fields, "2026-09-12", None)
        self.assertEqual(check_in, "2026-09-12")
        self.assertEqual(check_out, "2026-09-13")


class ImportCommandTests(TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.path = build_document(self._tmp.name)

    def test_dry_run_writes_nothing(self):
        run_import(self.path)
        self.assertEqual(Trip.objects.count(), 0)

    def test_apply_creates_trip_with_meals(self):
        run_import(self.path, apply=True)
        trip = Trip.objects.get()
        self.assertEqual(Meal.objects.filter(day__trip=trip).count(), 2)
        self.assertEqual(Day.objects.filter(trip=trip).count(), 1)

    def test_dry_run_previews_even_when_trip_exists(self):
        run_import(self.path, apply=True)
        # A preview of an already-imported document must still work, otherwise
        # you cannot re-check the mapping after fixing the parser.
        run_import(self.path)
        self.assertEqual(Trip.objects.count(), 1)

    def test_second_apply_is_blocked(self):
        run_import(self.path, apply=True)
        with self.assertRaises(CommandError):
            run_import(self.path, apply=True)
        self.assertEqual(Trip.objects.count(), 1)

    def test_replace_swaps_the_trip(self):
        run_import(self.path, apply=True)
        original = Trip.objects.get()
        run_import(self.path, apply=True, replace=True)
        self.assertEqual(Trip.objects.count(), 1)
        self.assertNotEqual(Trip.objects.get().pk, original.pk)
        # The old trip's days must cascade away, not linger.
        self.assertEqual(Day.objects.filter(trip_id=original.pk).count(), 0)

    def test_replace_reports_the_old_pk(self):
        run_import(self.path, apply=True)
        original_pk = Trip.objects.get().pk
        out = io.StringIO()
        with redirect_stdout(out):
            call_command("import_itinerary", str(self.path), apply=True, replace=True)
        # Django nulls an instance's pk once it is deleted, so this catches the
        # bug where the report said "Replaced Trip None".
        self.assertIn(f"Replaced Trip {original_pk}", out.getvalue())
