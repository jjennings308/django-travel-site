"""Tests for the Europe 2027 draft parser.

The real draft is a machine-generated document that is never committed, so
these tests build a miniature one with the same *shapes*: a Badge/Travelers
key table, day headings that carry a bracket badge instead of a "Day N"
number, bold sub-headings among plain body text, leg tables that give
durations rather than departure times, a Food Scene table, and a Key
Bookings list.

The properties worth pinning are the ones that would silently corrupt an
import rather than fail it: a badge that resolves to the wrong roster, a
weekday that disagrees with its own date, and a leg table quietly turned
into a TransportLeg with invented times.
"""

import io
import tempfile
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

from django.core.management import call_command
from django.test import TestCase

from apps.trips.docx_import import parse_document
from apps.trips.models import (
    BookingTask,
    Day,
    Lodging,
    Meal,
    Section,
    TransportLeg,
    Traveler,
    Trip,
)


def run_import(path, **kwargs):
    out, err = io.StringIO(), io.StringIO()
    with redirect_stdout(out), redirect_stderr(err):
        call_command("import_itinerary", str(path), **kwargs)


def docx_from(path):
    import docx

    return docx.Document(path)


def _bold(document, text):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    return paragraph


def _table(document, headers, rows):
    table = document.add_table(rows=1, cols=len(headers))
    for i, head in enumerate(headers):
        table.rows[0].cells[i].text = head
    for cells in rows:
        row = table.add_row()
        for i, value in enumerate(cells):
            row.cells[i].text = value
    return table


def build_europe_document(tmpdir, *, year=2027):
    """A miniature document in the Europe draft's shape."""
    import docx

    path = Path(tmpdir) / "europe.docx"
    document = docx.Document()

    _bold(document, "EUROPE 2027")
    document.add_paragraph("Bavaria · Windsor")
    document.add_paragraph("James · Regan (+Sara, Germany leg)")
    document.add_paragraph(f"September 17 – October 3, {year}")

    document.add_heading("Trip Overview", level=1)
    _table(
        document,
        ("Detail", "Info"),
        [
            ("Travelers (Germany)", "James, Regan, Sara"),
            ("Transport (Germany)", "Train-only throughout"),
            ("Oktoberfest", "Tuesday Sept 21"),
        ],
    )

    document.add_heading("Group Key & Hotels", level=1)
    _table(
        document,
        ("Badge", "Travelers", "Leg", "Notes"),
        [
            ("All 3", "James, Regan, Sara", "Germany", "Together"),
            ("Group of 2", "James, Regan", "England", "Sara flies home"),
        ],
    )
    _table(
        document,
        ("Hotel", "Location", "Dates", "Nights", "Key Features"),
        [
            ("Hessischer Hof", "Butzbach Old Town", "Sept 18–20", "2", "Historic centre"),
            ("TBD — Some Airbnb", "Augsburg", "Sept 20–25", "5", "Book immediately"),
        ],
    )

    document.add_heading("Germany Leg", level=1)
    # The last day of the walk is a Heading 1 away; the day walk must stop
    # there rather than at the "Germany Leg" heading, or England is lost.
    document.add_paragraph(f"Fri Sept 17  [All 3]  Overnight flight — PIT to FRA")
    _bold(document, "Departure")
    document.add_paragraph("Delta evening departure from Pittsburgh.")
    _table(
        document,
        ("Leg", "Mode", "Time", "Notes"),
        [("Butzbach → Frankfurt Hbf", "RE30", "~45 min", "Frequent service")],
    )
    _table(
        document,
        ("Restaurant / Venue", "Price", "Why it fits"),
        [("Ristorante Da Franco", "€€", "Wood-fired pizza"), ("Hotel breakfast", "incl.", "Buffet")],
    )
    _table(document, ("💡 Tip", "Olga is arranged personally by John"), [])
    document.add_paragraph("Sat Sept 18  [All 3]  Arrive FRA — Butzbach   Hotel 1: Hessischer Hof")
    _bold(document, "Morning — Arrive")
    document.add_paragraph("Land FRA ~7–9am, clear customs.")
    _bold(document, "Options — All Car-Free")
    _bold(document, "OPTION A (recommended): Hampton Court — train Windsor (~35 min).")
    # Deliberately not bold: the real draft only bolds option A, so matching
    # on formatting would silently drop every later option.
    document.add_paragraph("OPTION B: Oxford — train Windsor (~50 min). Bodleian Library.")
    document.add_paragraph("Sun Sept 19  [Group of 2]  Butzbach — Marburg   Hotel 1: Hessischer Hof")
    _bold(document, "Afternoon — Marburg")
    document.add_paragraph("Student guide at 3pm.")
    document.add_heading("Quick Reference", level=1)
    _table(
        document,
        ("Priority", "Item", "Deadline", "Notes"),
        [
            ("1 — CRITICAL", "Oktoberfest tent", "Jan 1, 2027", "Book early"),
            ("7 — High", "Hessischer Hof", "6+ months", "3 rooms"),
            ("13 — Medium", "Marburg student guide", "1–3 months", "Afternoon tour"),
        ],
    )
    _table(
        document,
        ("Hotel", "Address", "Phone", "Dates", "Rooms"),
        [("Hessischer Hof", "Farbgasse 1, Butzbach", "(+49) 123", "Sept 18–20", "3")],
    )

    document.save(path)
    return path


class EuropeParseTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = build_europe_document(self.tmp.name)
        self.parsed = parse_document(self.path)

    def test_document_is_recognised_as_europe_not_taos(self):
        # A Taos document has no day headings this parser recognises, so a
        # successful parse is itself the assertion that dispatch worked.
        self.assertEqual(len(self.parsed.days), 3)

    def test_trip_metadata_comes_from_the_title_block(self):
        self.assertEqual(self.parsed.trip["start_date"], "2027-09-17")
        self.assertEqual(self.parsed.trip["end_date"], "2027-10-03")
        self.assertIn("Bavaria", self.parsed.trip["name"])
        self.assertIn("EUROPE 2027", self.parsed.trip["name"])

    def test_travelers_come_from_the_group_key(self):
        self.assertEqual(
            [t["name"] for t in self.parsed.travelers], ["James", "Regan", "Sara"]
        )

    def test_badge_resolves_to_a_roster(self):
        day = self.parsed.days[0]
        self.assertEqual(day["badge"], "All 3")
        self.assertEqual(day["roster"], ["James", "Regan", "Sara"])

    def test_a_narrow_badge_resolves_to_a_subset(self):
        day = self.parsed.days[2]
        self.assertEqual(day["roster"], ["James", "Regan"])

    def test_unmatched_badge_warns_rather_than_guessing(self):
        parsed = parse_document(self._rewrite_badge("Group of 9"))
        self.assertEqual(parsed.days[0]["roster"], [])
        self.assertTrue(any("Group of 9" in w for w in parsed.warnings))

    def _rewrite_badge(self, badge):
        import docx

        document = docx.Document(self.path)
        for paragraph in document.paragraphs:
            if "Fri Sept 17" in paragraph.text:
                for run in paragraph.runs:
                    if "All 3" in run.text:
                        run.text = run.text.replace("All 3", badge)
        out = Path(self.tmp.name) / "rogue.docx"
        document.save(out)
        return out

    def test_weekday_disagreeing_with_the_date_warns(self):
        # 2027-09-17 is a Friday. Shifting the year makes every date wrong,
        # which is why the parser cross-checks rather than trusting.
        parsed = parse_document(build_europe_document(self.tmp.name, year=2028))
        self.assertEqual(parsed.trip["start_date"], "2028-09-17")
        self.assertTrue(
            any("says Fri" in w and "2028-09-17" in w for w in parsed.warnings),
            parsed.warnings,
        )

    def test_the_title_year_does_not_override_the_date_range(self):
        # The document's own title reads "EUROPE 2027" while the dates say
        # 2028. The date range is the one that has to win, or the whole trip
        # is filed a year out with no error raised.
        parsed = parse_document(build_europe_document(self.tmp.name, year=2028))
        self.assertEqual(parsed.trip["start_date"], "2028-09-17")
        self.assertEqual(parsed.trip["end_date"], "2028-10-03")

    def test_day_numbers_are_assigned_in_document_order(self):
        self.assertEqual([d["day_number"] for d in self.parsed.days], [1, 2, 3])

    def test_a_split_arrow_badge_still_means_the_whole_group_that_day(self):
        # "All 3 → splits" means the day starts as everybody and only splits
        # afterwards, so the day's own plan still covers everyone.
        document = docx_from(self.path)
        for paragraph in document.paragraphs:
            if "Sun Sept 19" in paragraph.text:
                for run in paragraph.runs:
                    run.text = run.text.replace("Group of 2", "All 3 → splits")
        path = Path(self.tmp.name) / "split.docx"
        document.save(path)
        self.assertEqual(parse_document(path).days[2]["roster"], ["James", "Regan", "Sara"])

    def test_bold_subheading_becomes_a_phase_and_body_becomes_its_summary(self):
        sections = self.parsed.days[0]["sections"]
        phase = next(s for s in sections if s["section_type"] == "phase")
        self.assertEqual(phase["title"], "Departure")
        self.assertIn("Delta evening", phase["content"]["summary"])

    def test_leg_table_is_kept_as_a_table_not_an_invented_leg(self):
        # A duration is not a departure time. Turning "~45 min" into a
        # TransportLeg would silently shift the leg by hours.
        sections = self.parsed.days[0]["sections"]
        table = next(s for s in sections if s["section_type"] == "logistics_table")
        self.assertEqual(table["content"]["columns"], ["Leg", "Mode", "Time", "Notes"])
        self.assertEqual(self.parsed.transport, [])

    def test_food_scene_table_becomes_meals(self):
        meals = self.parsed.days[0]["meals"]
        self.assertEqual([m["name"] for m in meals], ["Ristorante Da Franco", "Hotel breakfast"])

    def test_euro_price_falls_back_to_a_dollar_band_and_keeps_the_symbol(self):
        meal = self.parsed.days[0]["meals"][0]
        self.assertEqual(meal["price_range"], "$$")
        self.assertEqual(meal["reservation_notes"], "€€")

    def test_one_column_callout_table_becomes_a_callout(self):
        callout = next(
            s for s in self.parsed.days[0]["sections"] if s["section_type"] == "callout"
        )
        self.assertEqual(callout["title"], "💡 Tip")
        self.assertIn("Olga", callout["content"]["text"])

    def test_option_lines_become_a_choose_your_adventure(self):
        # The second OPTION line is not bold in the real draft, so matching on
        # formatting alone would silently drop it.
        cya = next(
            s
            for s in self.parsed.days[1]["sections"]
            if s["section_type"] == "choose_your_adventure"
        )
        options = cya["content"]["options"]
        self.assertEqual([o["letter"] for o in options], ["A", "B"])
        self.assertEqual(options[0]["title"], "Hampton Court")
        self.assertEqual(options[1]["title"], "Oxford")

    def test_hotel_legs_become_lodging_but_a_tbd_hotel_does_not(self):
        names = [record["name"] for record in self.parsed.lodging]
        self.assertEqual(names, ["Hessischer Hof"])
        self.assertEqual(str(self.parsed.lodging[0]["check_in"]), "2027-09-18")
        self.assertTrue(any("not booked yet" in w for w in self.parsed.warnings))

    def test_hotel_phone_becomes_a_contact(self):
        self.assertEqual(self.parsed.contacts[0]["phone"], "(+49) 123")

    def test_key_bookings_become_prioritised_tasks(self):
        tasks = self.parsed.booking_tasks
        self.assertEqual(len(tasks), 3)
        self.assertEqual(
            [t["priority"] for t in tasks],
            [BookingTask.Priority.CRITICAL, BookingTask.Priority.HIGH, BookingTask.Priority.MEDIUM],
        )
        self.assertIn("Jan 1, 2027", tasks[0]["notes"])

    def test_counts_include_tasks(self):
        self.assertEqual(self.parsed.counts()["tasks"], 3)


class TipBoxTests(TestCase):
    """Two-column tip boxes become callouts, however many tips they stack."""

    def handle(self, rows):
        from apps.trips.docx_europe import _handle_table

        day = {"sections": [], "meals": []}
        _handle_table(day, rows, warn=lambda *a, **k: None)
        return day["sections"]

    def test_a_two_tip_box_becomes_two_callouts_not_a_table(self):
        # The real draft stacks "📅 Book" over "💡 Tip" in one table. Read as a
        # table, the first tip became the column headings.
        sections = self.handle(
            [["⚠ CRITICAL", "Tent reservations open January 2027."], ["💡 Tip", "Arrive mid-morning."]]
        )
        self.assertEqual([s["section_type"] for s in sections], ["callout", "callout"])
        self.assertEqual([s["title"] for s in sections], ["⚠ CRITICAL", "💡 Tip"])
        self.assertEqual([s["content"]["tone"] for s in sections], ["warning", "info"])
        self.assertEqual(sections[0]["content"]["text"], "Tent reservations open January 2027.")
        self.assertEqual([s["order"] for s in sections], [1, 2])

    def test_a_note_label_is_a_tip(self):
        sections = self.handle([["💡 Tip", "Take the 9pm train."], ["⚠ Note", "Last train is 11pm."]])
        self.assertEqual([s["section_type"] for s in sections], ["callout", "callout"])

    def test_a_real_two_column_table_stays_a_table(self):
        sections = self.handle([["Detail", "Info"], ["Check-in", "3 PM"]])
        self.assertEqual([s["section_type"] for s in sections], ["logistics_table"])
        self.assertEqual(sections[0]["content"]["columns"], ["Detail", "Info"])

    def test_a_table_that_is_only_partly_tips_is_not_split(self):
        sections = self.handle([["💡 Tip", "Book early."], ["Check-in", "3 PM"]])
        self.assertEqual([s["section_type"] for s in sections], ["logistics_table"])


class EuropeImportTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = build_europe_document(self.tmp.name)

    def test_dry_run_writes_nothing(self):
        run_import(self.path)
        self.assertEqual(Trip.objects.count(), 0)

    def test_apply_creates_the_trip_with_rosters_and_tasks(self):
        run_import(self.path, apply=True)
        trip = Trip.objects.get()
        self.assertEqual(trip.travelers.count(), 3)
        self.assertEqual(trip.days.count(), 3)
        self.assertEqual(trip.booking_tasks.count(), 3)

    def test_a_whole_group_day_is_left_blank_rather_than_stored(self):
        # Days 1 and 2 carry an "All 3" badge. Blank already means the whole
        # group, so storing three names would be redundant data that also
        # makes is_split claim the group narrowed when it did not.
        run_import(self.path, apply=True)
        for day in Day.objects.filter(day_number__in=(1, 2)):
            self.assertEqual(list(day.travelers.all()), [])
            self.assertFalse(day.is_split)
            self.assertEqual(day.roster.count(), 3)

    def test_only_the_narrow_day_gets_a_roster(self):
        run_import(self.path, apply=True)
        self.assertEqual(
            list(
                Day.objects.filter(travelers__isnull=False)
                .values_list("day_number", flat=True)
                .distinct()
            ),
            [3],
        )
        day = Day.objects.get(day_number=3)
        self.assertTrue(day.is_split)
        self.assertEqual(day.roster_summary, "James, Regan")

    def test_sections_and_meals_are_written(self):
        run_import(self.path, apply=True)
        # phase(Departure) + logistics_table(leg) + callout
        # + phase(Morning) + phase(Options) + choose_your_adventure
        # + phase(Afternoon)
        self.assertEqual(Section.objects.count(), 7)
        self.assertEqual(Meal.objects.count(), 2)

    def test_lodging_and_contacts_are_written(self):
        run_import(self.path, apply=True)
        trip = Trip.objects.get()
        self.assertEqual(Lodging.objects.filter(trip=trip).count(), 1)
        self.assertEqual(trip.contacts.count(), 1)

    def test_leg_tables_do_not_become_transport_rows(self):
        run_import(self.path, apply=True)
        self.assertEqual(TransportLeg.objects.count(), 0)

    def test_a_second_import_is_refused_as_a_duplicate(self):
        from django.core.management.base import CommandError

        run_import(self.path, apply=True)
        with self.assertRaises(CommandError):
            run_import(self.path, apply=True)

    def test_replace_reimports_cleanly(self):
        run_import(self.path, apply=True)
        first = Trip.objects.get().pk
        run_import(self.path, apply=True, replace=True)
        self.assertEqual(Trip.objects.count(), 1)
        self.assertNotEqual(Trip.objects.get().pk, first)
        self.assertEqual(Day.objects.count(), 3)
